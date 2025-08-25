"""
Text parsers for MD, TXT, and DOCX formats.
"""

import logging
from typing import Dict, Any, Optional
from pathlib import Path
import base64

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

from .base import BaseParser, ParsedContent
from .factory import register_parser

logger = logging.getLogger(__name__)


@register_parser('md')
class MarkdownParser(BaseParser):
    """
    Parser for Markdown files.
    
    Handles .md and .markdown files, preserving the markdown formatting
    while extracting metadata from frontmatter if present.
    """
    
    def can_parse(self, source: Dict[str, Any]) -> bool:
        """Check if this parser can handle the given source."""
        return source.get('type', '').lower() in ['md', 'markdown']
    
    async def parse(self, source: Dict[str, Any]) -> ParsedContent:
        """
        Parse Markdown content.
        
        Args:
            source: Source dictionary with markdown content
            
        Returns:
            ParsedContent object with markdown content
        """
        self.logger.info("📝 Starting Markdown parsing")
        
        try:
            # Get content
            content = self._get_content(source)
            
            # Extract frontmatter and content
            frontmatter, main_content = self._extract_frontmatter(content)
            
            # Extract title from frontmatter or content
            title = self._extract_title(frontmatter, main_content)
            
            # Create metadata
            metadata = self._extract_metadata(source)
            if frontmatter:
                metadata.update(frontmatter)
            metadata['format'] = 'markdown'
            
            # Create origin information
            origin = self._create_origin_info(source)
            
            # Detect language
            language = self._detect_language(main_content)
            
            parsed_content = ParsedContent(
                content_md=self._clean_content(main_content),
                title=title,
                metadata=metadata,
                origin=origin,
                language=language
            )
            
            self.logger.info(f"✅ Markdown parsed successfully: {parsed_content.get_summary()}")
            return parsed_content
            
        except Exception as e:
            self.logger.error(f"❌ Markdown parsing failed: {str(e)}")
            raise ValueError(f"Failed to parse Markdown: {str(e)}")
    
    def _get_content(self, source: Dict[str, Any]) -> str:
        """Get content from source."""
        if 'content' in source:
            return source['content']
        elif 'file_path' in source:
            file_path = Path(source['file_path'])
            return file_path.read_text(encoding='utf-8')
        else:
            raise ValueError("Markdown source must have either 'content' or 'file_path'")
    
    def _extract_frontmatter(self, content: str) -> tuple[Dict[str, Any], str]:
        """
        Extract YAML frontmatter from markdown content.
        
        Returns:
            Tuple of (frontmatter_dict, main_content)
        """
        frontmatter = {}
        main_content = content
        
        if content.startswith('---\n'):
            parts = content.split('---\n', 2)
            if len(parts) >= 3:
                try:
                    import yaml
                    frontmatter = yaml.safe_load(parts[1]) or {}
                    main_content = parts[2]
                except ImportError:
                    # YAML not available, skip frontmatter parsing
                    pass
                except Exception as e:
                    self.logger.warning(f"Failed to parse frontmatter: {e}")
        
        return frontmatter, main_content
    
    def _extract_title(self, frontmatter: Dict[str, Any], content: str) -> Optional[str]:
        """Extract title from frontmatter or content."""
        # Try frontmatter first
        if 'title' in frontmatter:
            return str(frontmatter['title']).strip()
        
        # Try first heading
        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if line.startswith('# '):
                return line[2:].strip()
        
        return None


@register_parser('txt')
class TextParser(BaseParser):
    """
    Parser for plain text files.
    
    Handles .txt files and other plain text formats.
    """
    
    def can_parse(self, source: Dict[str, Any]) -> bool:
        """Check if this parser can handle the given source."""
        return source.get('type', '').lower() in ['txt', 'text']
    
    async def parse(self, source: Dict[str, Any]) -> ParsedContent:
        """
        Parse plain text content.
        
        Args:
            source: Source dictionary with text content
            
        Returns:
            ParsedContent object with text content
        """
        self.logger.info("📄 Starting text parsing")
        
        try:
            # Get content
            content = self._get_content(source)
            
            # Extract title (first non-empty line)
            title = self._extract_title(content)
            
            # Create metadata
            metadata = self._extract_metadata(source)
            metadata['format'] = 'text'
            
            # Create origin information
            origin = self._create_origin_info(source)
            
            # Detect language
            language = self._detect_language(content)
            
            parsed_content = ParsedContent(
                content_md=self._clean_content(content),
                title=title,
                metadata=metadata,
                origin=origin,
                language=language
            )
            
            self.logger.info(f"✅ Text parsed successfully: {parsed_content.get_summary()}")
            return parsed_content
            
        except Exception as e:
            self.logger.error(f"❌ Text parsing failed: {str(e)}")
            raise ValueError(f"Failed to parse text: {str(e)}")
    
    def _get_content(self, source: Dict[str, Any]) -> str:
        """Get content from source."""
        if 'content' in source:
            return source['content']
        elif 'file_path' in source:
            file_path = Path(source['file_path'])
            return file_path.read_text(encoding='utf-8')
        else:
            raise ValueError("Text source must have either 'content' or 'file_path'")
    
    def _extract_title(self, content: str) -> Optional[str]:
        """Extract title from first non-empty line."""
        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if line and len(line) > 3 and len(line) < 200:
                return line
        return None


@register_parser('docx')
class DocxParser(BaseParser):
    """
    Parser for Microsoft Word DOCX files.
    
    Extracts text content and basic metadata from DOCX files.
    """
    
    def __init__(self):
        """Initialize DOCX parser."""
        super().__init__()
        
        if DocxDocument is None:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install it with: pip install python-docx"
            )
    
    def can_parse(self, source: Dict[str, Any]) -> bool:
        """Check if this parser can handle the given source."""
        return source.get('type', '').lower() == 'docx'
    
    async def parse(self, source: Dict[str, Any]) -> ParsedContent:
        """
        Parse DOCX content.
        
        Args:
            source: Source dictionary with DOCX content
            
        Returns:
            ParsedContent object with extracted text
        """
        self.logger.info("📄 Starting DOCX parsing")
        
        try:
            # Open document
            doc = self._open_docx_document(source)
            
            # Extract text content
            content_md = self._extract_text_content(doc)
            
            # Extract metadata
            metadata = self._extract_docx_metadata(doc, source)
            
            # Extract title
            title = self._extract_title(doc, metadata)
            
            # Create origin information
            origin = self._create_origin_info(source)
            
            # Detect language
            language = self._detect_language(content_md)
            
            parsed_content = ParsedContent(
                content_md=self._clean_content(content_md),
                title=title,
                metadata=metadata,
                origin=origin,
                language=language
            )
            
            self.logger.info(f"✅ DOCX parsed successfully: {parsed_content.get_summary()}")
            return parsed_content
            
        except Exception as e:
            self.logger.error(f"❌ DOCX parsing failed: {str(e)}")
            raise ValueError(f"Failed to parse DOCX: {str(e)}")
    
    def _open_docx_document(self, source: Dict[str, Any]) -> DocxDocument:
        """Open DOCX document from source."""
        if 'file_path' in source:
            file_path = Path(source['file_path'])
            if not file_path.exists():
                raise ValueError(f"DOCX file not found: {file_path}")
            return DocxDocument(str(file_path))
        
        elif 'content' in source:
            # Handle base64 encoded content
            content = source['content']
            if isinstance(content, str):
                try:
                    docx_bytes = base64.b64decode(content)
                except Exception as e:
                    raise ValueError(f"Invalid base64 DOCX content: {e}")
            else:
                docx_bytes = content
            
            import io
            return DocxDocument(io.BytesIO(docx_bytes))
        
        else:
            raise ValueError("DOCX source must have either 'file_path' or 'content'")
    
    def _extract_text_content(self, doc: DocxDocument) -> str:
        """Extract text content from DOCX document."""
        content_parts = []
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                # Check if it's a heading (basic heuristic)
                if paragraph.style.name.startswith('Heading'):
                    level = paragraph.style.name.replace('Heading ', '')
                    if level.isdigit():
                        level = int(level)
                        content_parts.append(f"{'#' * level} {text}")
                    else:
                        content_parts.append(f"## {text}")
                else:
                    content_parts.append(text)
        
        return '\n\n'.join(content_parts)
    
    def _extract_docx_metadata(self, doc: DocxDocument, source: Dict[str, Any]) -> Dict[str, Any]:
        """Extract metadata from DOCX document."""
        metadata = self._extract_metadata(source)
        
        # Document properties
        core_props = doc.core_properties
        
        metadata.update({
            'docx_title': core_props.title or '',
            'docx_author': core_props.author or '',
            'docx_subject': core_props.subject or '',
            'docx_created': str(core_props.created) if core_props.created else '',
            'docx_modified': str(core_props.modified) if core_props.modified else '',
            'docx_last_modified_by': core_props.last_modified_by or '',
            'paragraph_count': len(doc.paragraphs),
            'format': 'docx'
        })
        
        return metadata
    
    def _extract_title(self, doc: DocxDocument, metadata: Dict[str, Any]) -> Optional[str]:
        """Extract title from DOCX document."""
        # Try document properties first
        docx_title = metadata.get('docx_title', '').strip()
        if docx_title:
            return docx_title
        
        # Try first heading or paragraph
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text and len(text) > 3 and len(text) < 200:
                return text
        
        return None
